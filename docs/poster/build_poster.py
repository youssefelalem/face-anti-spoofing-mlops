"""Remplit le template officiel ICAIS'26 (A0) avec le contenu du projet.

Les chiffres sont lus dans reports/*.json (aucune valeur recopiee a la main).

Usage (Python avec python-pptx) :
    python docs/poster/build_poster.py
Sortie : docs/poster/ICAIS26_Poster_FaceAntiSpoofing.pptx
"""

from __future__ import annotations

import copy
import json
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "ICAIS26_Poster_Template_A0.pptx"
OUTPUT = HERE / "ICAIS26_Poster_FaceAntiSpoofing.pptx"
FIG = ROOT / "reports" / "figures"
REP = ROOT / "reports"


def load(name: str) -> dict:
    return json.loads((REP / name).read_text(encoding="utf-8"))


FINAL = load("final_metrics.json")["test"]          # {nuaa|webcam: {modele: metriques}}
V1 = load("test_metrics_v1_subject_protocol.json")   # CNN v1 : ecart validation / test
DQ = {c["check"]: c for c in load("data_quality_report.json")["checks"]}
APP = load("app_benchmark.json")                     # taille du modele, latence mesuree
CNN, LBP, FUS = "CNN (MobileNetV2)", "LBP + SVM", "Fusion"


def webcam_stats() -> dict:
    import csv

    rows = list(csv.DictReader((ROOT / "data" / "processed" / "webcam_manifest.csv").open(encoding="utf-8")))
    train = [r for r in rows if r["split"] != "test"]
    return {"train_sessions": sorted({r["session"] for r in train}), "n_train": len(train),
            "n_people": len({r["subject"] for r in rows})}


WEBCAM = webcam_stats()
# Lissage temporel reellement utilise par l'application (lu dans le code de la page)
SMOOTHING = int(re.search(r"const SMOOTHING = (\d+)",
                          (ROOT / "app" / "static" / "app.js").read_text(encoding="utf-8")).group(1))

# --- Identite (a verifier par l'etudiant) -------------------------------------
TITLE = ("Face Anti-Spoofing under Session Shift: Fusing a MobileNetV2 CNN with LBP Texture "
         "in a Deployable MLOps Pipeline")
AUTHORS = "Youssef EL ALEM¹  ·  Supervised by Pr. Hafsa BENADDI¹"
AFFILIATION = ("¹ Master Data Science & Big Data (DSBD), Faculty of Sciences Ben M'Sik, "
               "Hassan II University of Casablanca, Morocco")

NAVY = RGBColor(0x0B, 0x1E, 0x3F)
INK = RGBColor(0x16, 0x23, 0x3B)
MUTED = RGBColor(0x5C, 0x65, 0x76)
BLUE = RGBColor(0x1F, 0x5F, 0xA8)
ICE = RGBColor(0xE8, 0xEE, 0xF8)


def pct(v: float) -> str:
    return f"{v * 100:.2f} %"


# --- Outils texte : on garde le style du template ---------------------------------
def shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def set_lines(sh, lines, size=None):
    """Remplace le texte d'une zone en conservant la mise en forme du premier run.

    lines : liste de lignes ; une ligne est une str ou une liste de (texte, gras).
    Les lignes sont separees par des sauts de ligne <a:br/> dans un seul paragraphe,
    comme dans le template.
    """
    p = sh.text_frame.paragraphs[0]._p
    first_r = p.find(qn("a:r"))
    base_rpr = copy.deepcopy(first_r.find(qn("a:rPr")))
    if size is not None:
        base_rpr.set("sz", str(int(size * 100)))
    for el in list(p):
        if el.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            p.remove(el)
    for extra in sh.text_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    end = p.find(qn("a:endParaRPr"))

    def add(el):
        if end is not None:
            end.addprevious(el)
        else:
            p.append(el)

    for i, line in enumerate(lines):
        if i:
            br = p.makeelement(qn("a:br"), {})
            br.append(copy.deepcopy(base_rpr))
            add(br)
        segs = [(line, None)] if isinstance(line, str) else line
        for text, bold in segs:
            r = p.makeelement(qn("a:r"), {})
            rpr = copy.deepcopy(base_rpr)
            if bold is not None:
                rpr.set("b", "1" if bold else "0")
            r.append(rpr)
            t = r.makeelement(qn("a:t"), {})
            t.text = text
            if text != text.strip():
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            r.append(t)
            add(r)


def delete(slide, *names):
    for n in names:
        el = shape(slide, n)._element
        el.getparent().remove(el)


def picture(slide, path, x, y, w, h, align="center"):
    """Insere une image dans la boite (x, y, w, h) en cm, en conservant ses proportions."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else x
    py = y + (h - ph) / 2
    return slide.shapes.add_picture(str(path), Cm(px), Cm(py), Cm(pw), Cm(ph))


def text_box(slide, x, y, w, h, runs, size=18, color=MUTED, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    for text, b, c, s in runs:
        r = p.add_run()
        r.text = text
        r.font.name = "Arial"
        r.font.size = Pt(s or size)
        r.font.bold = bold if b is None else b
        r.font.color.rgb = c or color
    return tb


def sub_heading(slide, x, y, w, text):
    text_box(slide, x, y, w, 1.2, [(text, True, BLUE, 18)])


def stat(slide, x, y, w, value, label):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(4.4))
    card.adjustments[0] = 0.12
    card.fill.solid()
    card.fill.fore_color.rgb = ICE
    card.line.fill.background()
    card.shadow.inherit = False
    text_box(slide, x, y + 0.6, w, 2.0, [(value, True, NAVY, 34)], align=PP_ALIGN.CENTER)
    text_box(slide, x + 0.3, y + 2.85, w - 0.6, 1.2, [(label, False, MUTED, 14)], align=PP_ALIGN.CENTER)


def m(test: str, model: str, key: str) -> float:
    r = FINAL[test][model]
    return r[key] if key in r else r["at_validation_threshold"][key]


def main() -> None:
    prs = Presentation(str(TEMPLATE))
    s = prs.slides[0]
    has_webcam = "webcam" in FINAL
    wt = "webcam" if has_webcam else "nuaa"
    n_nuaa = FINAL["nuaa"][CNN]["n"]
    lum = DQ["biais_luminance_moyenne"]["detail"]["auc_feature_seule"]
    sat = DQ["biais_saturation_moyenne"]["detail"]["auc_feature_seule"]

    # ---- En-tete --------------------------------------------------------------
    set_lines(shape(s, "TextBox 12"), [TITLE], size=44)
    set_lines(shape(s, "TextBox 14"), [AUTHORS])
    set_lines(shape(s, "TextBox 15"), [AFFILIATION])
    # Pas de PAPER ID pour les PFA (consigne de Pr. BENADDI, 2026-10-16) : cartouche supprime
    delete(s, "Rounded Rectangle 16", "TextBox 17")

    # ---- 1. Introduction ------------------------------------------------------
    set_lines(shape(s, "TextBox 24"), [
        "Face recognition unlocks phones and opens doors, yet a printed photo or a phone screen "
        "showing an authorised user is often enough to fool it: a presentation attack.",
        "",
        [("Face anti-spoofing (FAS)", True), (" decides whether the face in front of the camera is "
         "live (bona fide) or an artefact, before identity is checked.", False)],
        "",
        [("Research gap: ", True), ("deep models report near-perfect scores on standard splits, but "
         "these splits share the recording conditions of the training data. How well does a "
         "lightweight CNN hold up when the capture session changes, as it always does in deployment?",
          False)],
    ], size=21)

    # ---- 2. Objectives --------------------------------------------------------
    set_lines(shape(s, "TextBox 31"), [
        [("Main objective: ", True), ("detect print and screen attacks in real time from one RGB image "
         "on an ordinary laptop.", False)],
        [("Question: ", True), ("does a fine-tuned MobileNetV2 generalise to unseen capture sessions, "
         "and can texture descriptors make it more robust?", False)],
        "• Audit training data with a formal data contract",
        "• Compare CNN, LBP + SVM and their fusion (ISO/IEC 30107-3)",
        "• Deploy the detector as a containerised web application",
    ], size=21)

    # ---- 3. Data --------------------------------------------------------------
    delete(s, "TextBox 39", "TextBox 40")
    shape(s, "Rounded Rectangle 38").line.fill.background()
    # Uniquement le visage de l'auteur (aucune image NUAA ni visage d'un tiers sur le poster)
    picture(s, REP / "samples_webcam.png", 4.2, 56.0, 22.2, 7.7)
    lines = [
        [("NUAA Imposter DB [1]: ", True), ("12,614 face crops, 15 subjects, 3 sessions, print attacks. "
         f"Train: sessions 1-2 · test: session 3 ({n_nuaa:,} images, never seen).", False)],
    ]
    if has_webcam:
        nw = FINAL["webcam"][CNN]["n"]
        train_sessions = ", ".join(WEBCAM["train_sessions"])
        lines.append([("Own webcam set (shown above): ", True),
                      ("live faces and phone-screen replays recorded with the demo laptop; sessions "
                       f"{train_sessions} for training ({WEBCAM['n_train']} img), a new person in a new "
                       f"session for test ({nw} img).", False)])
    lines.append([("Data-contract audit: ", True),
                  (f"brightness alone separates the classes (AUC {lum:.2f}), "
                   f"and so does colour saturation (AUC {sat:.2f}): a shortcut risk.", False)])
    set_lines(shape(s, "TextBox 41"), lines, size=16)

    # ---- 4. Keywords + demo ---------------------------------------------------
    for name, kw in [("TextBox 49", "Anti-Spoofing"), ("TextBox 51", "Transfer Learning"),
                     ("TextBox 53", "CNN"), ("TextBox 55", "LBP texture"), ("TextBox 57", "MLOps")]:
        set_lines(shape(s, name), [kw])
    sub_heading(s, 3.9, 82.0, 23.0, "LIVE DEMO APPLICATION")
    shot = REP / "app_screenshot.png"
    if shot.is_file():
        picture(s, shot, 3.9, 83.6, 23.0, 20.5)
    text_box(s, 3.9, 104.6, 23.0, 6.0, [
        ("Webcam → face detection (OpenCV) → CNN (ONNX Runtime) + LBP-SVM → fused verdict in ",
         None, None, 18),
        (f"{APP['latency_ms_median']:.0f} ms", True, NAVY, 18),
        (f" per frame on a laptop CPU, smoothed over {SMOOTHING} frames. FastAPI service in a Docker image "
         f"({APP['onnx_mb']:.1f} MB model, no GPU).", None, None, 18),
    ])

    # ---- 5. Methodology : pipeline MLOps (comment) + modele de detection (quoi) ----
    step_shapes = ["TextBox 64"]
    for rr, ov, tbs in [(65, 66, (67, 68, 69)), (70, 71, (72, 73, 74)), (75, 76, (77, 78, 79)),
                        (80, 81, (82, 83, 84))]:
        step_shapes += [f"Rounded Rectangle {rr}", f"Oval {ov}"] + [f"TextBox {t}" for t in tbs]
    delete(s, *step_shapes)          # remplaces par le schema du pipeline, plus informatif
    picture(s, FIG / "mlops_pipeline.png", 30.4, 25.0, 23.4, 12.0)
    text_box(s, 30.6, 37.1, 23.0, 1.0, [
        (f"{APP['n_quality_checks']} data-contract checks · {APP['n_tests']} automated tests · dataset "
         "SHA-256 · MLflow · GitHub Actions · ONNX · Docker", None, MUTED, 12)])
    sub_heading(s, 30.6, 38.4, 23.0, "DETECTION MODEL")
    delete(s, "TextBox 86", "TextBox 87")
    shape(s, "Rounded Rectangle 85").line.fill.background()
    picture(s, FIG / "fusion_architecture.png", 30.9, 40.2, 22.3, 8.1)
    text_box(s, 30.6, 48.9, 23.0, 1.0, [
        ("CNN: frozen head 3 epochs (lr 1e-3), then full fine-tuning (lr 1e-4), early stopping. "
         "Fusion = mean of calibrated probabilities.", None, MUTED, 12)])

    # ---- 6. Results -----------------------------------------------------------
    delete(s, "TextBox 95", "TextBox 96", "TextBox 98", "TextBox 99")
    for n in ("Rounded Rectangle 94", "Rounded Rectangle 97"):
        shape(s, n).line.fill.background()
    picture(s, FIG / "eer_comparison.png", 30.8, 55.4, 22.6, 12.3)
    picture(s, FIG / "roc_nuaa.png", 30.8, 69.1, 11.2, 10.9)
    picture(s, FIG / f"roc_{wt}.png" if has_webcam else FIG / "validation_gap.png", 42.2, 69.1, 11.2, 10.9)
    where = "a new person in a new webcam session" if has_webcam else "the unseen NUAA session"
    set_lines(shape(s, "TextBox 100"), [
        [("KEY RESULT", True)],
        [(f"On {where}, fusing texture with the CNN cuts the EER from {m(wt, CNN, 'eer') * 100:.1f} % "
          f"(CNN alone) to {m(wt, FUS, 'eer') * 100:.1f} %. LBP-SVM alone scores "
          f"{m(wt, LBP, 'eer') * 100:.1f} %.", False)],
    ], size=17)
    shape(s, "TextBox 100").height = Cm(5.0)
    sw, gap, x0 = 5.35, 0.45, 30.6
    for i, (v, lab) in enumerate([(m(wt, FUS, "auc"), "AUC"), (m(wt, FUS, "ACER"), "ACER"),
                                  (m(wt, FUS, "APCER"), "APCER"), (m(wt, FUS, "BPCER"), "BPCER")]):
        stat(s, x0 + i * (sw + gap), 85.2, sw, f"{v:.3f}" if lab == "AUC" else pct(v), lab)
    text_box(s, 30.6, 89.9, 23.0, 1.0, [
        (f"Fusion, {'own webcam test (session s2)' if has_webcam else 'NUAA test'}, threshold fixed on "
         "validation", None, MUTED, 12)])
    picture(s, FIG / "validation_gap.png", 30.6, 91.4, 23.0, 14.6)
    text_box(s, 30.6, 106.6, 23.0, 3.6, [
        ("APCER: attacks accepted · BPCER: live faces rejected · ACER = (APCER + BPCER) / 2 "
         "· EER: point where APCER = BPCER.", None, MUTED, 12)])

    # ---- 7. Key findings ------------------------------------------------------
    findings = [
        ("TextBox 109", "TextBox 110", "Internal validation is misleading",
         f"CNN: AUC {V1['validation']['auc']:.4f} on validation, {V1['auc']:.3f} on the unseen session."),
        ("TextBox 113", "TextBox 114", "The data contract caught the shortcut",
         f"Brightness alone separates training classes (AUC {lum:.2f}): a cue absent in session 3."),
        ("TextBox 117", "TextBox 118", "Texture complements deep features",
         f"LBP-SVM AUC {m('nuaa', LBP, 'auc'):.3f} vs CNN {m('nuaa', CNN, 'auc'):.3f} on NUAA session 3."),
    ]
    for t_name, d_name, title, desc in findings:
        set_lines(shape(s, t_name), [title], size=20)
        set_lines(shape(s, d_name), [desc], size=17)

    # ---- 8. Discussion --------------------------------------------------------
    set_lines(shape(s, "TextBox 125"), [
        [("Why the CNN drops. ", True), ("Session-3 prints are of higher quality with natural colours, "
         "while training prints are pale and bright. The CNN partly learned these cues; removing them "
         "(greyscale, photometric normalisation) did not fix it, so the shift is deeper than colour.",
          False)],
        "",
        [("Why fusion helps. ", True), ("LBP describes micro-texture lost by printing or screens; the CNN "
         "captures higher-level cues. Their errors are partly independent.", False)],
        "",
        [("Limitations. ", True), ("the CNN is unstable: adding webcam data of a single person lowered "
         "its generalisation to other people. NUAA test results were inspected during design "
         f"(optimistic); the webcam set is small ({WEBCAM['n_people']} people); no 3D masks.", False)],
    ], size=21)

    # ---- 9. Conclusion (carte agrandie : la place liberee sous les references) ----
    shift = 20.0
    for n in ("Rounded Rectangle 126", "Rectangle 127"):
        shape(s, n).height = shape(s, n).height + Cm(shift)
    set_lines(shape(s, "TextBox 132"), [
        "• Near-perfect validation scores do not guarantee robustness to a new capture session.",
        "• Auditing data and fusing CNN with texture features gives a more reliable detector.",
        "• The full pipeline is reproducible and runs live on a laptop CPU.",
        "",
        [("Future work: ", True), ("larger multi-device data (e.g. CelebA-Spoof), domain generalisation "
         "training, and monitoring of the deployed model.", False)],
    ], size=23)
    take = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(57.4), Cm(72.0), Cm(22.9), Cm(13.5))
    take.adjustments[0] = 0.08
    take.fill.solid()
    take.fill.fore_color.rgb = ICE
    take.line.fill.background()
    take.shadow.inherit = False
    text_box(s, 58.6, 73.6, 20.5, 1.6, [("TAKE-HOME MESSAGE", True, BLUE, 22)])
    text_box(s, 58.6, 76.6, 20.5, 8.5, [
        ("Validate anti-spoofing models on a new capture session, never only on data recorded like the "
         "training set: that is where shortcuts hide and where deployment fails.", True, NAVY, 28)])

    # ---- 10. References : references et remerciements uniquement ----------------
    for n in ("Rounded Rectangle 133", "Rectangle 134", "Oval 135", "TextBox 136", "TextBox 137",
              "Connector 138", "TextBox 139"):
        shape(s, n).top = shape(s, n).top + Cm(shift)
    for n in ("Rounded Rectangle 133", "Rectangle 134"):
        shape(s, n).height = shape(s, n).height - Cm(shift)
    set_lines(shape(s, "TextBox 139"), [
        "[1] X. Tan et al. Face liveness detection from a single image with sparse low rank bilinear "
        "discriminative model. ECCV 2010.",
        "[2] M. Sandler et al. MobileNetV2: Inverted residuals and linear bottlenecks. CVPR 2018.",
        "[3] J. Määttä, A. Hadid, M. Pietikäinen. Face spoofing detection from single images "
        "using micro-texture analysis. IJCB 2011.",
        "[4] ISO/IEC 30107-3:2017. Biometric presentation attack detection - Testing and reporting.",
        "",
        [("Acknowledgements: ", True), ("Pr. Hafsa BENADDI for her supervision; the NUAA team for sharing "
         "the database for research.", False)],
    ], size=24)

    prs.save(str(OUTPUT))
    print("Poster :", OUTPUT)


if __name__ == "__main__":
    main()
